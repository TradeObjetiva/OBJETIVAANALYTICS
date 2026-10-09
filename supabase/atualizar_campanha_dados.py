import openpyxl
import json
import os

EXCEL_PATH = r"C:\Users\TRADE\Documents\arquivo campanha\CAMPANHA (1) (1).xlsx"

CAMPANHA_PERIODOS = [
    {"id": 1, "name": "27/08 A 03/09", "start": "2026-08-27", "end": "2026-09-03"},
    {"id": 2, "name": "04/09 A 11/09", "start": "2026-09-04", "end": "2026-09-11"},
    {"id": 3, "name": "12/09 A 19/09", "start": "2026-09-12", "end": "2026-09-19"},
    {"id": 4, "name": "21/09 A 26/09", "start": "2026-09-21", "end": "2026-09-26"},
    {"id": 5, "name": "28/09 A 30/09", "start": "2026-09-28", "end": "2026-09-30"},
    {"id": 6, "name": "01/10 A 07/10", "start": "2026-10-01", "end": "2026-10-07"},
    {"id": 7, "name": "08/10 A 14/10", "start": "2026-10-08", "end": "2026-10-14"},
    {"id": 8, "name": "15/10 A 21/10", "start": "2026-10-15", "end": "2026-10-21"},
    {"id": 9, "name": "22/10 A 31/10", "start": "2026-10-22", "end": "2026-10-31"},
    {"id": 10, "name": "01/11 A 07/11", "start": "2026-11-01", "end": "2026-11-07"},
    {"id": 11, "name": "08/11 A 14/11", "start": "2026-11-08", "end": "2026-11-14"},
    {"id": 12, "name": "15/11 A 21/11", "start": "2026-11-15", "end": "2026-11-21"},
    {"id": 13, "name": "22/11 A 30/11", "start": "2026-11-22", "end": "2026-11-30"},
    {"id": 14, "name": "01/12 A 07/12", "start": "2026-12-01", "end": "2026-12-07"},
    {"id": 15, "name": "08/12 A 14/12", "start": "2026-12-08", "end": "2026-12-14"},
    {"id": 16, "name": "15/12 A 19/12", "start": "2026-12-15", "end": "2026-12-19"}
]

# Colunas dos períodos 1 a 5 no Excel (base 1)
# P1: Assid(4), Efet(5), Pont(6), Ilha(7), MeiaIlha(8), Ponta(9), MeiaPonta(10), Tot(11)
# P2: 12..19
# P3: 20..27
# P4: 28..35
# P5: 36..43
PERIOD_OFFSETS = [
    (1, 4),
    (2, 12),
    (3, 20),
    (4, 28),
    (5, 36)
]

def clean_num(val):
    if val is None or val == "":
        return 0
    try:
        return int(float(str(val).replace(",", ".")))
    except:
        return 0

def run():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Ativos']

    colaboradores = []
    
    for r in range(3, ws.max_row + 1):
        nome_raw = ws.cell(r, 1).value
        if not nome_raw or not str(nome_raw).strip():
            continue
        nome = str(nome_raw).strip().upper()
        projeto = str(ws.cell(r, 2).value or "GERAL").strip().upper()
        equipe = str(ws.cell(r, 3).value or "GERAL").strip().upper()
        
        periodos_list = []
        pontos_extras_total = 0
        
        # Mapear os 16 períodos
        for p_info in CAMPANHA_PERIODOS:
            pid = p_info["id"]
            pname = p_info["name"]
            
            # Se for período 1..5, extrair da planilha
            if pid <= 5:
                start_c = 4 + (pid - 1) * 8
                assid = clean_num(ws.cell(r, start_c).value)
                efet = clean_num(ws.cell(r, start_c + 1).value)
                pont = clean_num(ws.cell(r, start_c + 2).value)
                ilha = clean_num(ws.cell(r, start_c + 3).value)
                meia_ilha = clean_num(ws.cell(r, start_c + 4).value)
                ponta = clean_num(ws.cell(r, start_c + 5).value)
                meia_ponta = clean_num(ws.cell(r, start_c + 6).value)
                
                pe_periodo = ilha + meia_ilha + ponta + meia_ponta
                pontos_extras_total += pe_periodo
                
                tot_calc = assid + pont + pe_periodo
                
                periodos_list.append({
                    "period_id": pid,
                    "period_name": pname,
                    "pontos_assiduidade": assid,
                    "efetividade": efet if efet > 0 else (100 if assid >= 0 else 50),
                    "pontos_pontualidade_total": pont,
                    "dias_pontuais": 0,
                    "dias_trabalhados": 6,
                    "faltas": 1 if assid < 0 else 0,
                    "ilha": ilha,
                    "meia_ilha": meia_ilha,
                    "ponta": ponta,
                    "meia_ponta": meia_ponta,
                    "total_periodo": tot_calc,
                    "ganhou_bonus": False
                })
            else:
                # Períodos futuros (6 a 16) - Prontos e zerados
                periodos_list.append({
                    "period_id": pid,
                    "period_name": pname,
                    "pontos_assiduidade": 0,
                    "efetividade": 100,
                    "pontos_pontualidade_total": 0,
                    "dias_pontuais": 0,
                    "dias_trabalhados": 6,
                    "faltas": 0,
                    "ilha": 0,
                    "meia_ilha": 0,
                    "ponta": 0,
                    "meia_ponta": 0,
                    "total_periodo": 0,
                    "ganhou_bonus": False
                })
        
        colaboradores.append({
            "nome": nome,
            "projeto": projeto,
            "equipe": equipe,
            "cargo": "PROMOTOR",
            "contrato": "ATIVO",
            "status": "OK",
            "sem_contrato": False,
            "pontos_extras": pontos_extras_total,
            "periodos": periodos_list
        })
        
    print(f"Sucesso: {len(colaboradores)} promotores processados.")
    com_pontos = [c for c in colaboradores if c["pontos_extras"] > 0]
    print(f"Promotores com pontos extras nos períodos preenchidos: {len(com_pontos)}")
    pts_total = sum(c["pontos_extras"] for c in colaboradores)
    print(f"Total de pontos extras: {pts_total}")
    
    # Salvar em public/subsystems/sistema-campanha/campanha_base_data.js
    out_js = "// Baseline da Campanha atualizada a partir de C:\\Users\\TRADE\\Documents\\arquivo campanha\\CAMPANHA (1) (1).xlsx\n"
    out_js += "window.CAMPANHA_BASE_DATA = " + json.dumps(colaboradores, ensure_ascii=False, indent=2) + ";\n"
    
    path_pub = r"public\subsystems\sistema-campanha\campanha_base_data.js"
    with open(path_pub, "w", encoding="utf-8") as f:
        f.write(out_js)
    print(f"Arquivo atualizado: {path_pub}")
    
    path_auto = r"AUTOMAÇÃO PARA CAMPANHA\static\campanha_base_data.js"
    if os.path.exists(os.path.dirname(path_auto)):
        with open(path_auto, "w", encoding="utf-8") as f:
            f.write(out_js)
        print(f"Arquivo atualizado: {path_auto}")

if __name__ == "__main__":
    run()

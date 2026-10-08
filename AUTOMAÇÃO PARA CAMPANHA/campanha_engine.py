import openpyxl
import os
import re
from datetime import datetime, timedelta

def parse_time_str(val):
    """Parses time strings like '08:00:00', '08:00', or datetime.time objects into seconds from midnight."""
    if val is None:
        return None
    if isinstance(val, (datetime,)):
        return val.hour * 3600 + val.minute * 60 + val.second
    if hasattr(val, 'hour') and hasattr(val, 'minute'):
        return val.hour * 3600 + val.minute * 60 + getattr(val, 'second', 0)
    
    s = str(val).strip()
    if not s or s.lower() in ['none', '-', '', 'nan']:
        return None
    
    is_neg = False
    if s.startswith('-'):
        is_neg = True
        s = s[1:].strip()
    
    parts = s.split(':')
    try:
        if len(parts) == 3:
            h, m, sec = int(parts[0]), int(parts[1]), int(float(parts[2]))
        elif len(parts) == 2:
            h, m, sec = int(parts[0]), int(parts[1]), 0
        else:
            return None
        total = h * 3600 + m * 60 + sec
        return -total if is_neg else total
    except (ValueError, TypeError):
        return None

def is_justified_absence(obs_str, dt_str):
    """
    Checks if an absence is justified.
    Returns (True, reason) if justified, (False, reason) if counted as an unjustified absence.
    """
    if dt_str in ['07/09/2026', '12/10/2026', '02/11/2026', '15/11/2026']:
        return True, f"Feriado Nacional ({dt_str})"
    
    if not obs_str:
        return False, "Falta sem registro / sem justificativa"
    
    obs_clean = str(obs_str).upper()
    
    justified_keywords = [
        'ATESTADO', 'MEDICO', 'MÉDICO', 'DECLARAÇÃO', 'DECLARACAO',
        'EXAME', 'TRATAMENTO', 'CONSULTA', 'FÉRIAS', 'FERIAS',
        'FERIADO', 'PRESENTE EM LOJA', 'SEM APP', 'SEM CELULAR',
        'FOLGA', 'ABONADA'
    ]
    
    for kw in justified_keywords:
        if kw in obs_clean:
            return True, f"Justificado: {kw}"
            
    return False, f"Falta injustificada: {obs_str.strip()}"

class CampanhaProcessor:
    def __init__(self, camp_file, folha_file, tolerance_minutes=10, bonus_points=30, colaboradores_base=None):
        self.camp_file = camp_file
        self.folha_file = folha_file
        self.tolerance_minutes = tolerance_minutes
        self.tolerance_seconds = tolerance_minutes * 60
        self.bonus_points = bonus_points
        self.colaboradores_base = colaboradores_base
        
        # Calendário Completo da Campanha: 27/08/2026 a 19/12/2026
        self.periods = [
            {'id': 1, 'name': '27/08 A 03/09', 'start': '2026-08-27', 'end': '2026-09-03', 'col_assid': 4, 'col_efet': 5, 'col_pont': 6, 'col_ilha': 7, 'col_meia_ilha': 8, 'col_ponta': 9, 'col_meia_ponta': 10, 'col_tot': 11},
            {'id': 2, 'name': '04/09 A 11/09', 'start': '2026-09-04', 'end': '2026-09-11', 'col_assid': 12, 'col_efet': 13, 'col_pont': 14, 'col_ilha': 15, 'col_meia_ilha': 16, 'col_ponta': 17, 'col_meia_ponta': 18, 'col_tot': 19},
            {'id': 3, 'name': '12/09 A 19/09', 'start': '2026-09-12', 'end': '2026-09-19', 'col_assid': 20, 'col_efet': 21, 'col_pont': 22, 'col_ilha': 23, 'col_meia_ilha': 24, 'col_ponta': 25, 'col_meia_ponta': 26, 'col_tot': 27},
            {'id': 4, 'name': '21/09 A 26/09', 'start': '2026-09-21', 'end': '2026-09-26', 'col_assid': 28, 'col_efet': 29, 'col_pont': 30, 'col_ilha': 31, 'col_meia_ilha': 32, 'col_ponta': 33, 'col_meia_ponta': 34, 'col_tot': 35},
            {'id': 5, 'name': '28/09 A 30/09', 'start': '2026-09-28', 'end': '2026-09-30', 'col_assid': 36, 'col_efet': 37, 'col_pont': 38, 'col_ilha': 39, 'col_meia_ilha': 40, 'col_ponta': 41, 'col_meia_ponta': 42, 'col_tot': 43},
            {'id': 6, 'name': '01/10 A 07/10', 'start': '2026-10-01', 'end': '2026-10-07', 'col_assid': 44, 'col_efet': 45, 'col_pont': 46, 'col_ilha': 47, 'col_meia_ilha': 48, 'col_ponta': 49, 'col_meia_ponta': 50, 'col_tot': 51},
            {'id': 7, 'name': '08/10 A 14/10', 'start': '2026-10-08', 'end': '2026-10-14', 'col_assid': 52, 'col_efet': 53, 'col_pont': 54, 'col_ilha': 55, 'col_meia_ilha': 56, 'col_ponta': 57, 'col_meia_ponta': 58, 'col_tot': 59},
            {'id': 8, 'name': '15/10 A 21/10', 'start': '2026-10-15', 'end': '2026-10-21', 'col_assid': 60, 'col_efet': 61, 'col_pont': 62, 'col_ilha': 63, 'col_meia_ilha': 64, 'col_ponta': 65, 'col_meia_ponta': 66, 'col_tot': 67},
            {'id': 9, 'name': '22/10 A 31/10', 'start': '2026-10-22', 'end': '2026-10-31', 'col_assid': 68, 'col_efet': 69, 'col_pont': 70, 'col_ilha': 71, 'col_meia_ilha': 72, 'col_ponta': 73, 'col_meia_ponta': 74, 'col_tot': 75},
            {'id': 10, 'name': '01/11 A 07/11', 'start': '2026-11-01', 'end': '2026-11-07', 'col_assid': 76, 'col_efet': 77, 'col_pont': 78, 'col_ilha': 79, 'col_meia_ilha': 80, 'col_ponta': 81, 'col_meia_ponta': 82, 'col_tot': 83},
            {'id': 11, 'name': '08/11 A 14/11', 'start': '2026-11-08', 'end': '2026-11-14', 'col_assid': 84, 'col_efet': 85, 'col_pont': 86, 'col_ilha': 87, 'col_meia_ilha': 88, 'col_ponta': 89, 'col_meia_ponta': 90, 'col_tot': 91},
            {'id': 12, 'name': '15/11 A 21/11', 'start': '2026-11-15', 'end': '2026-11-21', 'col_assid': 92, 'col_efet': 93, 'col_pont': 94, 'col_ilha': 95, 'col_meia_ilha': 96, 'col_ponta': 97, 'col_meia_ponta': 98, 'col_tot': 99},
            {'id': 13, 'name': '22/11 A 30/11', 'start': '2026-11-22', 'end': '2026-11-30', 'col_assid': 100, 'col_efet': 101, 'col_pont': 102, 'col_ilha': 103, 'col_meia_ilha': 104, 'col_ponta': 105, 'col_meia_ponta': 106, 'col_tot': 107},
            {'id': 14, 'name': '01/12 A 07/12', 'start': '2026-12-01', 'end': '2026-12-07', 'col_assid': 108, 'col_efet': 109, 'col_pont': 110, 'col_ilha': 111, 'col_meia_ilha': 112, 'col_ponta': 113, 'col_meia_ponta': 114, 'col_tot': 115},
            {'id': 15, 'name': '08/12 A 14/12', 'start': '2026-12-08', 'end': '2026-12-14', 'col_assid': 116, 'col_efet': 117, 'col_pont': 118, 'col_ilha': 119, 'col_meia_ilha': 120, 'col_ponta': 121, 'col_meia_ponta': 122, 'col_tot': 123},
            {'id': 16, 'name': '15/12 A 19/12', 'start': '2026-12-15', 'end': '2026-12-19', 'col_assid': 124, 'col_efet': 125, 'col_pont': 126, 'col_ilha': 127, 'col_meia_ilha': 128, 'col_ponta': 129, 'col_meia_ponta': 130, 'col_tot': 131},
        ]
        
    def find_sheet(self, folha_wb, nome_promotor):
        """Finds matching sheet in folha workbook by exact or truncated name."""
        if not folha_wb or not nome_promotor:
            return None
        nome_clean = nome_promotor.strip().upper()
        # 1. Exact match
        for s in folha_wb.sheetnames:
            if s.strip().upper() == nome_clean:
                return s
        # 2. Excel 31 character truncation or prefix match
        for s in folha_wb.sheetnames:
            s_clean = s.strip().upper()
            if nome_clean.startswith(s_clean) or s_clean.startswith(nome_clean[:25]):
                return s
        return None

    def analyze_promotor(self, ws_folha, nome_promotor, faltas_assiduidade_datas=None):
        """
        Analisa assiduidade e pontualidade do promotor.
        Cruza DIRETAMENTE com a tabela de assiduidade oficial (faltas_assiduidade_datas).
        """
        faltas_datas_set = set(faltas_assiduidade_datas or [])
        
        # Extract contract information from header
        contrato_val = None
        is_sem_contrato = False
        
        if ws_folha:
            for r in range(1, 9):
                for c in range(1, ws_folha.max_column + 1):
                    v = ws_folha.cell(r, c).value
                    if v and 'contrato' in str(v).lower() and str(v).strip().endswith(':'):
                        nv = ws_folha.cell(r, c + 1).value
                        if nv:
                            contrato_val = str(nv).strip()
            if not contrato_val:
                c7 = ws_folha.cell(7, 10).value
                if c7:
                    contrato_val = str(c7).strip()
                    
            jornada_header = str(ws_folha.cell(1, 8).value or '') + ' ' + str(ws_folha.cell(2, 8).value or '')
            if not contrato_val or 'sem contrato' in contrato_val.lower() or 'sem jornada' in jornada_header.lower():
                is_sem_contrato = True
                if not contrato_val:
                    contrato_val = "Usuário sem contrato"

        # Parse daily records from sheet if sheet exists
        daily_records = []
        if ws_folha:
            for r in range(11, 45):
                dt_cell = ws_folha.cell(r, 1).value
                if not dt_cell:
                    continue
                
                p = str(dt_cell).strip().split('/')
                if len(p) != 3:
                    continue
                dt_iso = f"{p[2]}-{p[1]}-{p[0]}"
                dt_br = f"{p[0]}/{p[1]}/{p[2]}"
                
                daily_records.append({
                    'row': r,
                    'date_br': dt_br,
                    'date_iso': dt_iso,
                    'dia_sem': str(ws_folha.cell(r, 2).value or ''),
                    'exec_in': ws_folha.cell(r, 3).value,
                    'exec_int_out': ws_folha.cell(r, 4).value,
                    'exec_int_ret': ws_folha.cell(r, 5).value,
                    'exec_out': ws_folha.cell(r, 6).value,
                    'exec_int_time': ws_folha.cell(r, 7).value,
                    'plan_in': ws_folha.cell(r, 9).value,
                    'plan_out': ws_folha.cell(r, 10).value,
                    'plan_int_time': ws_folha.cell(r, 11).value,
                    'diff_time': ws_folha.cell(r, 13).value,
                    'obs': str(ws_folha.cell(r, 14).value or '')
                })

        # Group by periods and calculate stats
        period_results = []
        consecutive_punctual_periods = 0
        
        for p_idx, p in enumerate(self.periods):
            p_days = [d for d in daily_records if p['start'] <= d['date_iso'] <= p['end']]
            
            # Faltas registradas diretamente na tabela oficial de assiduidade neste período
            faltas_oficiais_no_periodo = [dt for dt in faltas_datas_set if p['start'] <= dt <= p['end']]
            
            faltas = 0
            dias_planejados = 0
            dias_trabalhados = 0
            dias_pontuais = 0
            dias_com_atraso = 0
            detalhes_dias = []
            
            # 1. Analisa dias com registro na folha
            dias_processados_iso = set()
            for d in p_days:
                dias_processados_iso.add(d['date_iso'])
                plan_in_sec = parse_time_str(d['plan_in'])
                exec_in_sec = parse_time_str(d['exec_in'])
                plan_int_sec = parse_time_str(d['plan_int_time'])
                exec_int_sec = parse_time_str(d['exec_int_time'])
                
                day_status = 'FOLGA'
                pontual_hoje = False
                motivo_falta = ''
                atraso_min = 0
                estouro_int_min = 0
                
                # VERIFICAÇÃO 1: RECONHECIMENTO DIRETO DA TABELA DE ASSIDUIDADE (tb_assiduidade)
                if d['date_iso'] in faltas_datas_set:
                    dias_planejados += 1
                    day_status = 'FALTA'
                    faltas += 1
                    motivo_falta = 'Falta registrada diretamente na Tabela de Assiduidade'
                elif plan_in_sec is not None and plan_in_sec > 0:
                    dias_planejados += 1
                    
                    if exec_in_sec is None:
                        is_just, justificativa = is_justified_absence(d['obs'], d['date_br'])
                        if is_just:
                            day_status = 'AUSENCIA_JUSTIFICADA'
                            motivo_falta = justificativa
                        else:
                            day_status = 'FALTA'
                            faltas += 1
                            motivo_falta = justificativa
                    else:
                        dias_trabalhados += 1
                        day_status = 'TRABALHADO'
                        
                        delay_start_sec = exec_in_sec - plan_in_sec
                        atraso_min = max(0, round(delay_start_sec / 60, 1))
                        is_in_pontual = (delay_start_sec <= self.tolerance_seconds)
                        
                        is_int_pontual = True
                        if plan_int_sec is not None and plan_int_sec > 0:
                            if exec_int_sec is not None:
                                excess_int_sec = exec_int_sec - plan_int_sec
                                if excess_int_sec > self.tolerance_seconds:
                                    is_int_pontual = False
                                    estouro_int_min = round(excess_int_sec / 60, 1)
                                    
                        if is_in_pontual and is_int_pontual:
                            pontual_hoje = True
                            dias_pontuais += 1
                        else:
                            dias_com_atraso += 1
                            
                detalhes_dias.append({
                    'date_br': d['date_br'],
                    'dia_sem': d['dia_sem'],
                    'status': day_status,
                    'exec_in': d['exec_in'],
                    'plan_in': d['plan_in'],
                    'exec_int': d['exec_int_time'],
                    'plan_int': d['plan_int_time'],
                    'pontual': pontual_hoje,
                    'atraso_min': atraso_min,
                    'estouro_int_min': estouro_int_min,
                    'obs': d['obs'],
                    'motivo_falta': motivo_falta
                })
                
            # 2. Inclui eventuais faltas da tabela de assiduidade não listadas na folha
            for dt_iso in faltas_oficiais_no_periodo:
                if dt_iso not in dias_processados_iso:
                    faltas += 1
                    dias_planejados += 1
                    detalhes_dias.append({
                        'date_br': dt_iso,
                        'dia_sem': '',
                        'status': 'FALTA',
                        'exec_in': None,
                        'plan_in': '08:00',
                        'exec_int': None,
                        'plan_int': None,
                        'pontual': False,
                        'atraso_min': 0,
                        'estouro_int_min': 0,
                        'obs': 'Falta registrada em tb_assiduidade',
                        'motivo_falta': 'Falta reconhecida via Tabela de Assiduidade'
                    })

            # Score calculation
            pontos_assiduidade = -50 * faltas if faltas > 0 else 0
            pontos_pontualidade_base = dias_pontuais
            
            is_period_perfect = (dias_trabalhados > 0 and dias_com_atraso == 0 and faltas == 0)
            if is_period_perfect:
                consecutive_punctual_periods += 1
            else:
                consecutive_punctual_periods = 0
                
            ganhou_bonus = False
            bonus_complemento = 0
            if consecutive_punctual_periods == 4:
                ganhou_bonus = True
                soma_3_anteriores = sum(period_results[i]['pontos_pontualidade_total'] for i in range(p_idx - 3, p_idx))
                target_total_4 = self.bonus_points
                bonus_complemento = max(0, target_total_4 - (soma_3_anteriores + pontos_pontualidade_base))
                
            pontos_pontualidade_total = pontos_pontualidade_base + bonus_complemento
            
            period_results.append({
                'period_id': p['id'],
                'period_name': p['name'],
                'faltas': faltas,
                'pontos_assiduidade': pontos_assiduidade,
                'dias_planejados': dias_planejados,
                'dias_trabalhados': dias_trabalhados,
                'dias_pontuais': dias_pontuais,
                'dias_com_atraso': dias_com_atraso,
                'is_period_perfect': is_period_perfect,
                'consecutive_count': consecutive_punctual_periods,
                'ganhou_bonus': ganhou_bonus,
                'bonus_complemento': bonus_complemento,
                'pontos_pontualidade_base': pontos_pontualidade_base,
                'pontos_pontualidade_total': pontos_pontualidade_total,
                'detalhes_dias': detalhes_dias
            })
            
        return {
            'nome': nome_promotor,
            'contrato': contrato_val or 'Contrato Padrão CLT',
            'sem_contrato': is_sem_contrato,
            'periodos': period_results
        }

    def process_all(self, faltas_assiduidade_map=None, campanha_salva_map=None):
        """
        Processa toda a folha de ponto e colaboradores.
        Ajuste INCREMENTAL: reconhece dados já salvos e adiciona novos.
        """
        faltas_map = faltas_assiduidade_map or {}
        salva_map = campanha_salva_map or {}
        
        wb_folha = None
        if self.folha_file and os.path.exists(self.folha_file):
            try:
                wb_folha = openpyxl.load_workbook(self.folha_file, data_only=True)
            except Exception as e:
                print(f"Erro ao abrir folha {self.folha_file}: {e}")

        # Lista de promotores prioritária
        lista_promotores = []
        if self.colaboradores_base and len(self.colaboradores_base) > 0:
            for idx, c in enumerate(self.colaboradores_base):
                nome = c.get('nome')
                if not nome:
                    continue
                lista_promotores.append({
                    'row': idx + 3,
                    'nome': str(nome).strip(),
                    'projeto': str(c.get('projeto') or 'GERAL').strip(),
                    'equipe': str(c.get('equipe') or c.get('cargo') or 'PROMOTOR').strip()
                })
        elif self.camp_file and os.path.exists(self.camp_file):
            wb_camp = openpyxl.load_workbook(self.camp_file, data_only=True)
            ws_camp = wb_camp['Ativos'] if 'Ativos' in wb_camp.sheetnames else wb_camp.active
            for r in range(3, ws_camp.max_row + 1):
                nome = ws_camp.cell(r, 1).value
                if not nome:
                    continue
                lista_promotores.append({
                    'row': r,
                    'nome': str(nome).strip(),
                    'projeto': str(ws_camp.cell(r, 2).value or 'GERAL').strip(),
                    'equipe': str(ws_camp.cell(r, 3).value or 'PROMOTOR').strip()
                })
        
        results = []
        for item in lista_promotores:
            r = item['row']
            nome = item['nome']
            projeto = item['projeto']
            equipe = item['equipe']
            nome_clean = nome.strip().upper()
            
            # Faltas oficiais deste colaborador
            faltas_datas = faltas_map.get(nome_clean, [])
            
            sheet_name = self.find_sheet(wb_folha, str(nome)) if wb_folha else None
            ws_folha = wb_folha[sheet_name] if (wb_folha and sheet_name) else None
            
            # Verifica se já temos dados salvos deste colaborador no Supabase
            saved_colab = salva_map.get(nome_clean)
            
            analysis = self.analyze_promotor(ws_folha, str(nome), faltas_assiduidade_datas=faltas_datas)
            
            # Se já havia dados salvos na campanha para este colaborador, mescla de forma incremental
            if saved_colab and saved_colab.get('periodos'):
                saved_periodos = saved_colab['periodos']
                for p_curr in analysis['periodos']:
                    # Procura se esse período já tinha dados salvos consolidados
                    matched_saved = next((sp for sp in saved_periodos if sp.get('period_id') == p_curr['period_id'] or sp.get('period_name') == p_curr['period_name']), None)
                    if matched_saved:
                        # Preserva dados já salvos caso a folha nova não possua registro para aquele ciclo
                        if p_curr['dias_trabalhados'] == 0 and matched_saved.get('dias_trabalhados', 0) > 0:
                            p_curr['dias_trabalhados'] = matched_saved.get('dias_trabalhados', 0)
                            p_curr['dias_pontuais'] = matched_saved.get('dias_pontuais', 0)
                            p_curr['pontos_pontualidade_total'] = matched_saved.get('pontos_pontualidade_total', 0)
                            p_curr['pontos_assiduidade'] = matched_saved.get('pontos_assiduidade', 0)
                            p_curr['faltas'] = matched_saved.get('faltas', 0)

            results.append({
                'row': r,
                'nome': nome,
                'projeto': projeto,
                'equipe': equipe,
                'status': 'OK' if sheet_name else ('COM_FALTAS_ASSIDUIDADE' if faltas_datas else 'SEM_FOLHA'),
                'contrato': analysis['contrato'],
                'sem_contrato': analysis['sem_contrato'],
                'sheet_name': sheet_name,
                'periodos': analysis['periodos'],
                'pontos_extras': saved_colab.get('pontos_extras', 0) if saved_colab else 0
            })
            
        return results

    def generate_updated_excel(self, output_path):
        """
        Gera o Excel atualizado preservando fórmulas e valores existentes.
        """
        if not self.camp_file or not os.path.exists(self.camp_file):
            return None
        
        wb = openpyxl.load_workbook(self.camp_file, data_only=False)
        ws = wb['Ativos'] if 'Ativos' in wb.sheetnames else wb.active
        
        wb_folha = openpyxl.load_workbook(self.folha_file, data_only=True) if (self.folha_file and os.path.exists(self.folha_file)) else None
        
        for r in range(3, ws.max_row + 1):
            nome_val = ws.cell(r, 1).value
            if not nome_val:
                continue
            nome = str(nome_val).strip()
            
            sheet_name = self.find_sheet(wb_folha, nome) if wb_folha else None
            if not sheet_name:
                continue
                
            analysis = self.analyze_promotor(wb_folha[sheet_name], nome)
            
            for p in self.periods:
                p_res = next((res for res in analysis['periodos'] if res['period_id'] == p['id']), None)
                if not p_res:
                    continue
                # Escreve Assiduidade e Pontualidade se a coluna existir
                if p['col_assid'] <= ws.max_column:
                    ws.cell(r, p['col_assid']).value = p_res['pontos_assiduidade']
                if p['col_pont'] <= ws.max_column:
                    ws.cell(r, p['col_pont']).value = p_res['pontos_pontualidade_total']
                    
        wb.save(output_path)
        return output_path

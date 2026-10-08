import os
import io
import json
import re
import pandas as pd
import openpyxl
from datetime import datetime

class PontoExtraEngine:
    PONTUACOES_BASE = {
        'ILHA': 50,
        'MEIA_ILHA': 25,
        'PONTA': 30,
        'MEIA_PONTA': 15,
        'NENHUM': 0
    }
    
    LABELS_TIPO = {
        'ILHA': 'Ilha (50 pts)',
        'MEIA_ILHA': 'Meia Ilha (25 pts)',
        'PONTA': 'Ponta de Gôndola (30 pts)',
        'MEIA_PONTA': 'Meia Ponta (15 pts)',
        'NENHUM': 'Nenhum / Rejeitado (0 pts)'
    }

    def __init__(self, base_dir=None, camp_file=None, pe_csv_file=None, storage_file=None):
        self.base_dir = base_dir or os.path.dirname(os.path.abspath(__file__))
        self.camp_file = camp_file or os.path.join(self.base_dir, "CAMPANHA (1).xlsx")
        self.pe_csv_file = pe_csv_file or os.path.join(self.base_dir, "RELATÓRIO_PONTO_EXTRA (4).csv")
        self.storage_file = storage_file or os.path.join(self.base_dir, "aprovacoes_pontos_extras.json")
        self.load_storage()

    def load_storage(self):
        if os.path.exists(self.storage_file):
            try:
                with open(self.storage_file, 'r', encoding='utf-8') as f:
                    self.aprovacoes = json.load(f)
            except Exception:
                self.aprovacoes = {}
        else:
            self.aprovacoes = {}

    def save_storage(self):
        with open(self.storage_file, 'w', encoding='utf-8') as f:
            json.dump(self.aprovacoes, f, ensure_ascii=False, indent=2)

    def get_campaign_promoters_and_periods(self):
        """Reads promoter names and periods directly from CAMPANHA (1).xlsx."""
        if not os.path.exists(self.camp_file):
            return [], []

        wb = openpyxl.load_workbook(self.camp_file, data_only=True)
        ws = wb['Ativos']

        # Periods in Row 1 & 2
        periods = []
        p_id = 1
        for col in range(4, ws.max_column + 1, 8):
            name_val = ws.cell(1, col).value
            if not name_val:
                name_val = f"Período {p_id}"
            periods.append({
                'id': p_id,
                'name': str(name_val).strip(),
                'col_start': col,
                'col_ilha': col + 3,
                'col_meia_ilha': col + 4,
                'col_ponta': col + 5,
                'col_meia_ponta': col + 6,
                'col_total': col + 7
            })
            p_id += 1

        promoters = []
        for r in range(3, ws.max_row + 1):
            val = ws.cell(r, 1).value
            if val:
                promoters.append({
                    'row': r,
                    'nome': str(val).strip(),
                    'projeto': str(ws.cell(r, 2).value or '').strip(),
                    'equipe': str(ws.cell(r, 3).value or '').strip()
                })

        return promoters, periods

    def match_promoter(self, csv_name, camp_promoters):
        """Matches CSV promoter name with campaign spreadsheet."""
        if not csv_name:
            return None
        clean_csv = re.sub(r'\s+', ' ', str(csv_name).strip().upper())
        
        # 1. Exact match
        for p in camp_promoters:
            if p['nome'].upper() == clean_csv:
                return p['nome']
                
        # 2. Prefix / substring match
        for p in camp_promoters:
            p_clean = p['nome'].upper()
            if clean_csv.startswith(p_clean[:20]) or p_clean.startswith(clean_csv[:20]):
                return p['nome']
                
        return None

    def calculate_points(self, tipo, criativo, quantidade=1):
        """Calculates points based on type, quantity and creative double bonus."""
        base = self.PONTUACOES_BASE.get(tipo, 0)
        multiplicador = 2 if criativo else 1
        return base * int(quantidade) * multiplicador

    def load_ponto_extra_records(self):
        """Parses the CSV and attaches current approval statuses."""
        if not os.path.exists(self.pe_csv_file):
            return [], []

        camp_promoters, periods = self.get_campaign_promoters_and_periods()

        # Read CSV with pandas
        try:
            df = pd.read_csv(self.pe_csv_file, sep=';', skiprows=1, encoding='utf-8', low_memory=False)
        except Exception:
            df = pd.read_csv(self.pe_csv_file, sep=';', skiprows=1, encoding='latin1', low_memory=False)

        # Standardize columns
        photo_cols = [c for c in df.columns if 'FOTO' in c.upper()]
        
        # Filter rows with Ponto Extra = SIM or with at least one photo
        has_pe_col = 'TEM PONTO EXTRA?' in df.columns
        if has_pe_col:
            is_sim = df['TEM PONTO EXTRA?'].astype(str).str.strip().str.upper() == 'SIM'
            has_photo = df[photo_cols].notna().any(axis=1) if photo_cols else False
            df_pe = df[is_sim | has_photo].copy()
        else:
            df_pe = df[df[photo_cols].notna().any(axis=1)].copy()

        records = []
        for _, row in df_pe.iterrows():
            task_id = str(row['ID DA TAREFA']).strip()
            data_exec = str(row.get('DATA E HORA DE EXECUÇÃO', '') or '').strip()
            pdv = str(row.get('PDV', '') or '').strip()
            promotor_csv = str(row.get('PROMOTOR', '') or '').strip()
            atividade = str(row.get('ATIVIDADE', '') or '').strip()

            # Extract photos
            fotos = []
            for pc in photo_cols:
                val = row.get(pc)
                if pd.notna(val) and str(val).strip().startswith('http'):
                    fotos.append(str(val).strip())

            # Match promoter in campaign
            matched_promotor = self.match_promoter(promotor_csv, camp_promoters)

            # Check saved approval
            aprov = self.aprovacoes.get(task_id, {})
            status = aprov.get('status', 'PENDENTE')
            tipo = aprov.get('tipo', 'ILHA') # default suggestion
            criativo = bool(aprov.get('criativo', False))
            quantidade = int(aprov.get('quantidade', 1))
            periodo = aprov.get('periodo', 'P5 (28/09 A 30/09)') # or auto-derived
            motivo = aprov.get('motivo', '')
            pontos = self.calculate_points(tipo, criativo, quantidade) if status == 'APROVADO' else 0

            records.append({
                'task_id': task_id,
                'data_exec': data_exec,
                'pdv': pdv,
                'promotor_csv': promotor_csv,
                'promotor_campanha': matched_promotor or promotor_csv,
                'matched': bool(matched_promotor),
                'atividade': atividade,
                'fotos': fotos,
                'total_fotos': len(fotos),
                'status': status,
                'tipo': tipo,
                'criativo': criativo,
                'quantidade': quantidade,
                'periodo': periodo,
                'motivo': motivo,
                'pontos': pontos,
                'updated_at': aprov.get('updated_at', None)
            })

        return records, periods

    def update_approval(self, task_id, data):
        """Updates a single approval decision."""
        task_id = str(task_id).strip()
        status = data.get('status', 'PENDENTE')
        tipo = data.get('tipo', 'ILHA')
        criativo = bool(data.get('criativo', False))
        quantidade = max(1, int(data.get('quantidade', 1)))
        periodo = data.get('periodo', 'P5 (28/09 A 30/09)')
        motivo = data.get('motivo', '')
        pontos = self.calculate_points(tipo, criativo, quantidade) if status == 'APROVADO' else 0

        self.aprovacoes[task_id] = {
            'status': status,
            'tipo': tipo,
            'criativo': criativo,
            'quantidade': quantidade,
            'periodo': periodo,
            'motivo': motivo,
            'pontos': pontos,
            'updated_at': datetime.now().isoformat()
        }
        self.save_storage()
        return self.aprovacoes[task_id]

    def update_batch(self, updates_list):
        """Batch update approvals."""
        for item in updates_list:
            task_id = str(item.get('task_id', '')).strip()
            if task_id:
                self.update_approval(task_id, item)
        self.save_storage()

    def sync_to_campaign_excel(self, output_excel_path):
        """
        Calculates total approved points for ILHA, MEIA ILHA, PONTA, MEIA PONTA
        for each promoter and period, and writes them into CAMPANHA_ATUALIZADA.xlsx.
        """
        records, periods = self.load_ponto_extra_records()
        
        # Aggregate approved points by (promotor_campanha, periodo, tipo)
        aggregated = {} # key: (promotor_nome, periodo_str) -> {'ILHA': pts, 'MEIA_ILHA': pts, 'PONTA': pts, 'MEIA_PONTA': pts}
        
        for r in records:
            if r['status'] == 'APROVADO':
                prom = r['promotor_campanha']
                per = r['periodo']
                tipo = r['tipo']
                pts = r['pontos']
                
                key = (prom, per)
                if key not in aggregated:
                    aggregated[key] = {'ILHA': 0, 'MEIA_ILHA': 0, 'PONTA': 0, 'MEIA_PONTA': 0}
                if tipo in aggregated[key]:
                    aggregated[key][tipo] += pts

        # Open workbook (preserving formulas)
        src_file = output_excel_path if os.path.exists(output_excel_path) else self.camp_file
        wb = openpyxl.load_workbook(src_file, data_only=False)
        ws = wb['Ativos']

        # Map periods in workbook
        camp_promoters, camp_periods = self.get_campaign_promoters_and_periods()

        # Build promoter name to row mapping
        name_to_row = {}
        for r in range(3, ws.max_row + 1):
            val = ws.cell(r, 1).value
            if val:
                name_to_row[str(val).strip().upper()] = r

        # Write aggregated points into sheet
        feed_summary = []
        for (prom_name, per_str), tipos_dict in aggregated.items():
            prom_clean = prom_name.upper()
            row_idx = name_to_row.get(prom_clean)
            if not row_idx:
                # Try fuzzy
                for k, v in name_to_row.items():
                    if prom_clean.startswith(k[:20]) or k.startswith(prom_clean[:20]):
                        row_idx = v
                        break
            if not row_idx:
                continue

            # Find matching period in camp_periods
            target_period = None
            for p in camp_periods:
                if p['name'] in per_str or per_str in p['name'] or f"P{p['id']}" in per_str:
                    target_period = p
                    break
            
            # Default to period 5 if not matched
            if not target_period and camp_periods:
                target_period = camp_periods[-1]

            if target_period:
                # Col positions
                c_ilha = target_period['col_ilha']
                c_milha = target_period['col_meia_ilha']
                c_ponta = target_period['col_ponta']
                c_mponta = target_period['col_meia_ponta']

                if tipos_dict['ILHA'] > 0:
                    ws.cell(row_idx, c_ilha).value = tipos_dict['ILHA']
                if tipos_dict['MEIA_ILHA'] > 0:
                    ws.cell(row_idx, c_milha).value = tipos_dict['MEIA_ILHA']
                if tipos_dict['PONTA'] > 0:
                    ws.cell(row_idx, c_ponta).value = tipos_dict['PONTA']
                if tipos_dict['MEIA_PONTA'] > 0:
                    ws.cell(row_idx, c_mponta).value = tipos_dict['MEIA_PONTA']

                feed_summary.append({
                    'promotor': prom_name,
                    'periodo': target_period['name'],
                    'ilha': tipos_dict['ILHA'],
                    'meia_ilha': tipos_dict['MEIA_ILHA'],
                    'ponta': tipos_dict['PONTA'],
                    'meia_ponta': tipos_dict['MEIA_PONTA'],
                    'total': sum(tipos_dict.values())
                })

        # Ensure SUM formula in TOTAL column for all periods and rows
        from openpyxl.utils import get_column_letter
        for r in range(3, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                continue
            for p in camp_periods:
                c_start = p['col_start']
                c_end = c_start + 6
                c_tot = p['col_total']
                l_start = get_column_letter(c_start)
                l_end = get_column_letter(c_end)
                ws.cell(r, c_tot).value = f"=SUM({l_start}{r}:{l_end}{r})"

        wb.save(output_excel_path)
        return feed_summary

    def get_summary_stats(self):
        """Returns consolidated statistics of all extra points."""
        records, periods = self.load_ponto_extra_records()
        total_recs = len(records)
        aprovados = [r for r in records if r['status'] == 'APROVADO']
        rejeitados = [r for r in records if r['status'] == 'REJEITADO']
        pendentes = [r for r in records if r['status'] == 'PENDENTE']

        total_pts = sum(r['pontos'] for r in aprovados)
        total_criativos = sum(1 for r in aprovados if r['criativo'])

        by_tipo = {'ILHA': 0, 'MEIA_ILHA': 0, 'PONTA': 0, 'MEIA_PONTA': 0}
        pts_by_tipo = {'ILHA': 0, 'MEIA_ILHA': 0, 'PONTA': 0, 'MEIA_PONTA': 0}
        for r in aprovados:
            t = r['tipo']
            if t in by_tipo:
                by_tipo[t] += r['quantidade']
                pts_by_tipo[t] += r['pontos']

        # Distinct promoters
        promoters_set = sorted(list(set(r['promotor_campanha'] for r in records)))

        return {
            'total_registros': total_recs,
            'total_aprovados': len(aprovados),
            'total_rejeitados': len(rejeitados),
            'total_pendentes': len(pendentes),
            'total_pontos': total_pts,
            'total_criativos': total_criativos,
            'by_tipo': by_tipo,
            'pts_by_tipo': pts_by_tipo,
            'promoters': promoters_set,
            'periods': [p['name'] for p in periods]
        }

    def export_excel_audit(self):
        """Generates a detailed Excel workbook of all ponto extra approvals."""
        records, periods = self.load_ponto_extra_records()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Curadoria Pontos Extras"

        headers = [
            "ID Tarefa", "Status", "Promotor (CSV)", "Promotor (Campanha)", "PDV",
            "Data/Hora", "Atividade", "Período", "Tipo Conquista", "Criativo (Dobro)",
            "Qtd", "Pontos Computados", "Motivo / Observação", "Foto 01", "Foto 02", "Foto 03"
        ]
        ws.append(headers)

        for r in records:
            fotos = r['fotos']
            ws.append([
                r['task_id'],
                r['status'],
                r['promotor_csv'],
                r['promotor_campanha'],
                r['pdv'],
                r['data_exec'],
                r['atividade'],
                r['periodo'],
                r['tipo'],
                "SIM" if r['criativo'] else "NÃO",
                r['quantidade'],
                r['pontos'],
                r['motivo'],
                fotos[0] if len(fotos) > 0 else "",
                fotos[1] if len(fotos) > 1 else "",
                fotos[2] if len(fotos) > 2 else ""
            ])

        # Style headers
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

